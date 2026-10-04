#!/usr/bin/env python3
"""
MEB çerçeve yıllık planlarını (OGM / TYMM ve DÖGM Excel dosyaları) ayrıştırır.

Girdi : tools/kaynaklar/ogm/<ders>/*.xlsx  ve  tools/kaynaklar/dogm/*.xlsx
Çıktı : tools/ara/planlar_ham.json

Her sayfa (sınıf) için 37 haftalık satırlar; birleştirilmiş (merge) hücrelerin
değerleri alt satırlara taşınır, "devam" bilgisi işaretlenir.
"""
import glob
import json
import os
import re
import unicodedata

import openpyxl

KOK = os.path.dirname(os.path.abspath(__file__))
KAYNAK = os.path.join(KOK, "kaynaklar")
ARA = os.path.join(KOK, "ara")


def norm(s):
    if s is None:
        return ""
    s = str(s).replace("\xa0", " ").replace("\r", "").replace("​", "")
    return unicodedata.normalize("NFC", s)


def temiz(s):
    s = norm(s)
    satirlar = [re.sub(r"[ \t]+", " ", l).strip() for l in s.split("\n")]
    out = []
    for l in satirlar:
        if l == "" and (not out or out[-1] == ""):
            continue
        out.append(l)
    while out and out[-1] == "":
        out.pop()
    while out and out[0] == "":
        out.pop(0)
    return "\n".join(out)


def tr_upper(s):
    return s.replace("i", "İ").replace("ı", "I").upper()


def baslik_anahtari(h):
    h = re.sub(r"\s+", " ", tr_upper(temiz(h)).replace("\n", " ")).strip()
    h = h.replace(" / ", " ").replace("/", " ")
    if h in ("AY", "MONTH"):
        return "ay"
    if h.startswith("HAFTA") or h == "WEEK":
        return "hafta"
    if "SAAT" in h or h == "HOUR":
        return "saat"
    if h.startswith("ÜNİTE İÇERİK"):
        return None
    if h.startswith("ÜNİTE") or h == "TEMA" or h.startswith("ÖĞRENME ALANI"):
        return "unite"
    if h == "UNIT":
        return "unite_no"
    if h == "THEME":
        return "unite"
    if h.startswith("KONU"):
        return "konu"
    if h.startswith("ÖĞRENME ÇIKTI") or h == "LEARNING OUTCOMES":
        return "cikti"
    if h == "KAZANIM" or h.startswith("KAZANIMLAR"):
        return "cikti"
    if h.startswith("SÜREÇ") or h.startswith("KAZANIM AÇIKLAMA") or h == "EXPLANATIONS":
        return "surec"
    if h.startswith("ÖLÇME") or h.startswith("MEASUREMENT"):
        return "olcme"
    if h.startswith("YÖNTEM") or h.startswith("METHODS"):
        return "yontem"
    if h.startswith("BELİRLİ GÜN") or h.startswith("NATIONAL DAYS") or h == "ETKİNLİK":
        return "gunler"
    if h.startswith("ANAHTAR KAVRAM"):
        return "kavramlar"
    if h.startswith("DESTEKLEME"):
        return "destekleme"
    if h.startswith("ZENGİNLEŞTİRME"):
        return "zenginlestirme"
    if h.startswith("FARKLILAŞTIRMA"):
        return "farklilastirma"
    return None


def birlesik_haritasi(ws):
    harita = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                harita[(r, c)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
    return harita


HAFTA_RE = re.compile(r"^\s*(\d{1,2})\s*\.?\s*HAFTA", re.I)


def sayfa_ayristir(ws):
    baslik = temiz(ws.cell(1, 1).value)
    hdr = None
    for r in range(1, 8):
        for c in range(1, min(ws.max_column, 32) + 1):
            v = tr_upper(temiz(ws.cell(r, c).value))
            if v.startswith("HAFTA") or v == "WEEK":
                hdr = r
                break
        if hdr:
            break
    if not hdr:
        return None
    sutun = {}
    for c in range(1, min(ws.max_column, 32) + 1):
        k = baslik_anahtari(ws.cell(hdr, c).value)
        if k and k not in sutun:
            sutun[k] = c
    if "hafta" not in sutun:
        return None
    mg = birlesik_haritasi(ws)
    # Çok satırlı başlıklar (ör. Arapça planları): eşlenmemiş sütunlar için üst satırlara bak
    eslenen = set(sutun.values())
    for c in range(1, min(ws.max_column, 32) + 1):
        if c in eslenen:
            continue
        for r in (hdr - 1, hdr - 2):
            if r < 1:
                continue
            ham = ws.cell(r, c).value
            yatay_ikinci = False
            if not ham and (r, c) in mg:
                r0, c0, r1, c1 = mg[(r, c)]
                ham = ws.cell(r0, c0).value
                yatay_ikinci = c0 != c
            metin = tr_upper(temiz(ham))
            if not metin:
                continue
            if "İÇERİK ÇERÇEVESİ" in metin and "ÜNİTE" in metin:
                k = "konu" if yatay_ikinci else "unite"
            elif metin.startswith("ÖĞRENME KANITLARI") or metin.startswith("ÖLÇME"):
                k = "olcme"
            elif yatay_ikinci:
                k = None
            else:
                k = baslik_anahtari(ham)
            if k and k not in sutun:
                sutun[k] = c
                eslenen.add(c)
            break

    def hucre(r, c):
        """Hücre değeri; dikey birleşimde üst hücrenin değeri (devam=True)."""
        if (r, c) in mg:
            r0, c0, r1, c1 = mg[(r, c)]
            if c0 != c:  # yatay birleşimin sağ parçası
                return "", False
            if r0 != r:
                return temiz(ws.cell(r0, c0).value), True
        return temiz(ws.cell(r, c).value), False

    haftalar, tatiller, cur = [], [], None
    for r in range(hdr + 1, ws.max_row + 1):
        a = temiz(ws.cell(r, 1).value)
        hafta_deger, hafta_devam = hucre(r, sutun["hafta"])
        au = tr_upper(a)
        if a and "TATİL" in au and not (hafta_deger and not hafta_devam):
            tatiller.append(re.sub(r"\s+", " ", a))
            cur = None
            continue
        m = HAFTA_RE.match(hafta_deger.replace("\n", " ")) if hafta_deger else None
        tarih_only = (not m and hafta_deger and not hafta_devam
                      and re.search(r"\d{1,2}\s*[-–/]\s*\d{0,2}\s*[A-Za-zÇĞİÖŞÜçğıöşü]", hafta_deger)
                      and re.search(r"(EYL|EKİM|KAS|ARA|OCA|ŞUB|MAR|NİS|MAY|HAZ)", tr_upper(hafta_deger)))
        if (m and not hafta_devam) or tarih_only:
            no = int(m.group(1)) if m else len(haftalar) + 1
            tarih = hafta_deger.split(":", 1)[1] if (m and ":" in hafta_deger) else re.sub(r"^\s*\d+\s*\.?\s*Hafta\s*", "", hafta_deger, flags=re.I)
            tarih = re.sub(r"\s+", " ", tarih.replace("\n", " ")).strip(" /")
            cur = {"hafta": no, "tarih": tarih, "satirlar": []}
            haftalar.append(cur)
        if cur is None:
            continue
        satir = {}
        for k, c in sutun.items():
            if k in ("ay", "hafta"):
                continue
            v, devam = hucre(r, c)
            if v:
                satir[k] = v
                if devam:
                    satir.setdefault("_devam", []).append(k)
        if any(k in satir for k in ("saat", "unite", "konu", "cikti", "surec", "olcme", "gunler")):
            cur["satirlar"].append(satir)
    return {"baslik": baslik, "sutunlar": list(sutun.keys()), "haftalar": haftalar, "tatiller": tatiller}


def main():
    os.makedirs(ARA, exist_ok=True)
    sonuc = []
    dosyalar = sorted(glob.glob(os.path.join(KAYNAK, "ogm", "*", "*.xlsx"))) + sorted(glob.glob(os.path.join(KAYNAK, "dogm", "*.xlsx")))
    for f in dosyalar:
        kaynak = "dogm" if os.sep + "dogm" + os.sep in f else "ogm"
        klasor = os.path.basename(os.path.dirname(f)) if kaynak == "ogm" else "dogm"
        wb = openpyxl.load_workbook(f, data_only=True)
        for ws in wb.worksheets:
            p = sayfa_ayristir(ws)
            if not p:
                print("ATLANDI (başlık yok):", os.path.basename(f), ws.title)
                continue
            p.update({"kaynak": kaynak, "klasor": klasor, "dosya": unicodedata.normalize("NFC", os.path.basename(f)), "sayfa": ws.title.strip()})
            sonuc.append(p)
            print(f"{kaynak:4} {klasor:24} {ws.title.strip()[:28]:28} hafta={len(p['haftalar']):2} | {p['baslik'][:70]}")
    with open(os.path.join(ARA, "planlar_ham.json"), "w", encoding="utf-8") as fh:
        json.dump(sonuc, fh, ensure_ascii=False, indent=1)
    print("Toplam sayfa:", len(sonuc))


if __name__ == "__main__":
    main()
